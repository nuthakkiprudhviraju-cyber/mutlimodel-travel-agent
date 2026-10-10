import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    DARK_BG = RGBColor(15, 23, 42)      # Slate 900
    CARD_BG = RGBColor(30, 41, 59)      # Slate 800
    BORDER_COL = RGBColor(51, 65, 85)   # Slate 700
    PRIMARY_BLUE = RGBColor(37, 99, 235)# Blue 600
    ACCENT_CYAN = RGBColor(56, 189, 248)# Sky 400
    TEXT_WHITE = RGBColor(248, 250, 252)# Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)# Slate 400
    SUCCESS_GREEN = RGBColor(34, 197, 94)
    ROSE_RED = RGBColor(244, 63, 94)
    AMBER_GOLD = RGBColor(245, 158, 11)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.color.rgb = DARK_BG
        return bg

    def add_header(slide, title_text, category_text="ROAMAGENT AI TRAVEL PLATFORM"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_CYAN
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(26)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_before = Pt(4)

    # 1. Team & Problem
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    add_header(s1, "Team Introduction & Problem Statement", "Slide 01 • Project Overview")

    c1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.1))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COL
    tf1 = c1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = Inches(0.4)
    tf1.margin_top = Inches(0.4)
    p = tf1.paragraphs[0]
    p.text = "👥 Project Team & Hackathon Mission"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    items_t1 = [
        ("Team Name:", "RoamAgent Innovators (Samartha Hackathon 2026)"),
        ("Project:", "RoamAgent - Autonomous Multi-Modal Travel & Logistics Planner"),
        ("Focus Area:", "Autonomous AI Agents, Geospatial Optimization & Last-Mile Transit"),
        ("Core Capability:", "Multi-modal routing (Air, Rail, Bus, Cab) with Doorstep Rapido pickup and live OSM discovery.")
    ]
    for lbl, val in items_t1:
        p_lbl = tf1.add_paragraph()
        p_lbl.text = f"• {lbl} "
        p_lbl.font.bold = True
        p_lbl.font.size = Pt(13)
        p_lbl.font.color.rgb = PRIMARY_BLUE
        p_lbl.space_before = Pt(12)
        p_val = tf1.add_paragraph()
        p_val.text = f"   {val}"
        p_val.font.size = Pt(12)
        p_val.font.color.rgb = TEXT_WHITE
        p_val.space_before = Pt(2)

    c2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.1))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COL
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = Inches(0.4)
    tf2.margin_top = Inches(0.4)
    p = tf2.paragraphs[0]
    p.text = "⚠️ The Problem Statement"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ROSE_RED

    problems = [
        ("Fragmented Booking Portals", "Travelers juggle 4+ disconnected apps for flights, trains, cabs, and hotels, causing massive friction and schedule gaps."),
        ("The 'First/Last Mile' Black Hole", "Intercity booking neglects the commute from doorstep to terminal/station, leading to missed departures and unpredictable transit times."),
        ("Blind Spot in Transit Infrastructure", "Lack of integrated awareness of real-time route detours, highway amenities, emergency medical access, and dining stops."),
        ("Uncontrolled Itinerary Drift", "Post-confirmation changes cause corporate travel policy chaos, unexpected expenses, and logistical disarray.")
    ]
    for title, desc in problems:
        pt = tf2.add_paragraph()
        pt.text = f"✕ {title}"
        pt.font.bold = True
        pt.font.size = Pt(13)
        pt.font.color.rgb = ROSE_RED
        pt.space_before = Pt(10)
        pd = tf2.add_paragraph()
        pd.text = f"   {desc}"
        pd.font.size = Pt(11.5)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_before = Pt(2)

    # 2. Aim & Objectives
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Aim & Strategic Objectives", "Slide 02 • Project Purpose")

    aim_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.3))
    aim_box.fill.solid()
    aim_box.fill.fore_color.rgb = CARD_BG
    aim_box.line.color.rgb = PRIMARY_BLUE
    aim_tf = aim_box.text_frame
    aim_tf.word_wrap = True
    aim_tf.margin_left = Inches(0.4)
    aim_tf.margin_top = Inches(0.2)
    p = aim_tf.paragraphs[0]
    p.text = "🎯 OUR CORE AIM"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p2 = aim_tf.add_paragraph()
    p2.text = "To build an autonomous, AI-driven corporate travel ecosystem that transforms natural-language travel intent into unified, end-to-end itineraries—seamlessly integrating highway routing, doorstep Rapido pickups, and destination amenities."
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(4)

    obj_cards = [
        ("1. Zero-Friction AI Parsing", "Utilize Gemini 2.5 Flash agents to extract dates, multi-city origins/destinations, and budget limits directly from free-form user prompts with sub-second execution.", PRIMARY_BLUE),
        ("2. Doorstep-to-Destination Unity", "Unify intercity transport (Air/Rail/Bus) with first/last mile Rapido pickup & drop points, mapping the real-time commute right onto interactive Leaflet charts.", SUCCESS_GREEN),
        ("3. Rich Contextual Geolocation", "Integrate OpenStreetMap & Overpass API to dynamically discover regional tourist sights, verified restaurants, and medical emergency centers along the route.", ACCENT_CYAN),
        ("4. Immutable Trip State & Lock", "Provide a seamless review cockpit with a multi-selection confirmation modal and an immutable locked state once booked to guarantee itinerary compliance.", AMBER_GOLD)
    ]
    left_positions = [Inches(0.8), Inches(3.8), Inches(6.8), Inches(9.8)]
    for i, (title, text, color) in enumerate(obj_cards):
        ob = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_positions[i], Inches(3.1), Inches(2.7), Inches(3.7))
        ob.fill.solid()
        ob.fill.fore_color.rgb = CARD_BG
        ob.line.color.rgb = BORDER_COL
        otf = ob.text_frame
        otf.word_wrap = True
        otf.margin_left = otf.margin_right = Inches(0.25)
        otf.margin_top = Inches(0.3)
        op1 = otf.paragraphs[0]
        op1.text = title
        op1.font.size = Pt(14)
        op1.font.bold = True
        op1.font.color.rgb = color
        op2 = otf.add_paragraph()
        op2.text = text
        op2.font.size = Pt(11)
        op2.font.color.rgb = TEXT_MUTED
        op2.space_before = Pt(10)

    # 3. Preview with User's Screenshot
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Preview of Our Platform & Journey Planner", "Slide 03 • System Interface")

    img_path = r"C:\Users\naveena\.gemini\antigravity\brain\fbcb660c-1257-4fea-a8a3-8eed78404588\.user_uploaded\media_1791666900884_69f35726.png"
    if os.path.exists(img_path):
        s3.shapes.add_picture(img_path, Inches(0.8), Inches(1.7), width=Inches(7.3))

    pv_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.4), Inches(1.7), Inches(4.1), Inches(5.1))
    pv_box.fill.solid()
    pv_box.fill.fore_color.rgb = CARD_BG
    pv_box.line.color.rgb = BORDER_COL
    pv_tf = pv_box.text_frame
    pv_tf.word_wrap = True
    pv_tf.margin_left = pv_tf.margin_right = Inches(0.3)
    pv_tf.margin_top = Inches(0.3)

    p = pv_tf.paragraphs[0]
    p.text = "🖥️ Intelligent Planning Interface"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    ui_points = [
        ("Stage 01: Natural Language Prompt", "Travelers describe journey in plain language. Gemini AI parses departure, arrival, deadline, and budget instantly."),
        ("Multi-Modal Transport Modes", "Instant toggling between Intercity Bus, Express Train, and Highway Cab/Taxi options."),
        ("Dynamic Parameter Customization", "Granular controls for departure dates, arrival deadlines (e.g. before 9:00 AM), max budget (₹1,500), and seating tiers."),
        ("Seamless Cockpit Navigation", "Left navigation bar with direct access to Itineraries, Disruption Simulator, Trip Details, Activity Log, and Profile.")
    ]
    for head, body in ui_points:
        p_h = pv_tf.add_paragraph()
        p_h.text = f"✔ {head}"
        p_h.font.bold = True
        p_h.font.size = Pt(12)
        p_h.font.color.rgb = SUCCESS_GREEN
        p_h.space_before = Pt(8)
        p_b = pv_tf.add_paragraph()
        p_b.text = f"   {body}"
        p_b.font.size = Pt(10.5)
        p_b.font.color.rgb = TEXT_MUTED
        p_b.space_before = Pt(2)

    # 4. Workflow
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "End-to-End User & System Workflow", "Slide 04 • Operational Flow")

    steps = [
        ("STEP 1", "Natural Language Prompt", "User enters prompt: 'I want to travel from Chennai to Hyderabad on the selected date, reach before 9 AM, prefer sleeper bus and keep budget below ₹1,500'.", PRIMARY_BLUE),
        ("STEP 2", "Multi-Modal Aggregation", "Backend generates flight, train, bus, and cab options. OSRM engine plots highway driving path while Overpass API queries local POIs.", ACCENT_CYAN),
        ("STEP 3", "Interactive Customization", "User selects preferred transport ticket, toggles Rapido doorstep bike/auto pickup & drop, and picks tourist sights & restaurants.", SUCCESS_GREEN),
        ("STEP 4", "Double-Check Review", "User clicks 'Confirm Details'. A verification modal appears displaying all selected tickets, cabs, and amenities for instant review.", AMBER_GOLD),
        ("STEP 5", "Booking & Immutable Lock", "Trip is saved to 'Trips' view. User can edit while in draft. Once user clicks 'Book Trip', system locks details against any further edits.", ROSE_RED)
    ]
    for i, (step_num, title, desc, color) in enumerate(steps):
        y_pos = Inches(1.6 + i * 1.05)
        step_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(1.8), Inches(0.9))
        step_box.fill.solid()
        step_box.fill.fore_color.rgb = color
        step_box.line.color.rgb = color
        stf = step_box.text_frame
        stf.margin_top = Inches(0.18)
        sp = stf.paragraphs[0]
        sp.text = step_num
        sp.font.bold = True
        sp.font.size = Pt(14)
        sp.font.color.rgb = TEXT_WHITE
        sp.alignment = PP_ALIGN.CENTER
        
        desc_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.8), y_pos, Inches(9.7), Inches(0.9))
        desc_box.fill.solid()
        desc_box.fill.fore_color.rgb = CARD_BG
        desc_box.line.color.rgb = BORDER_COL
        dtf = desc_box.text_frame
        dtf.word_wrap = True
        dtf.margin_left = Inches(0.3)
        dtf.margin_top = Inches(0.12)
        dp1 = dtf.paragraphs[0]
        dp1.text = title
        dp1.font.bold = True
        dp1.font.size = Pt(13)
        dp1.font.color.rgb = TEXT_WHITE
        dp2 = dtf.add_paragraph()
        dp2.text = desc
        dp2.font.size = Pt(11)
        dp2.font.color.rgb = TEXT_MUTED
        dp2.space_before = Pt(2)

    # 5. Tech Stack
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Full-Stack Technology Architecture", "Slide 05 • Tech Stack")

    tech_cols = [
        ("🎨 Frontend Tier", [
            ("React 19 & TypeScript", "Ultra-fast reactive UI with strict typed component architecture."),
            ("Vite 6 Bundler", "Sub-second HMR and optimized production bundle compilation."),
            ("Tailwind CSS v4", "Modern utility-first styling with zero-overflow pinned layouts."),
            ("Leaflet & React-Leaflet", "Interactive map rendering, highway polyline overlay & Rapido pins.")
        ], ACCENT_CYAN, Inches(0.8)),
        
        ("⚡ Backend Tier", [
            ("Python 3.13 & FastAPI", "High-performance asynchronous REST API framework."),
            ("SQLAlchemy 2.0 ORM", "Type-safe database transactions and lifecycle modeling."),
            ("Neon PostgreSQL", "Serverless cloud PostgreSQL for resilient trip persistence."),
            ("Pydantic v2", "Strict schema validation and error payload serialization.")
        ], PRIMARY_BLUE, Inches(4.8)),

        ("🧠 AI, Geo & External APIs", [
            ("Google Gemini 2.5 Flash", "Autonomous prompt parameter extraction & itinerary generation."),
            ("OSRM (Open Source Routing)", "Real-time highway driving coordinates, duration, and road geometry."),
            ("Overpass OSM Engine", "Live geographical discovery of attractions, food & hospitals."),
            ("Curated Mock Engine", "Instant zero-downtime offline fallback for zero API failures.")
        ], SUCCESS_GREEN, Inches(8.8))
    ]
    for col_title, items, col_color, x_pos in tech_cols:
        col_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.7), Inches(3.7), Inches(5.1))
        col_box.fill.solid()
        col_box.fill.fore_color.rgb = CARD_BG
        col_box.line.color.rgb = BORDER_COL
        ctf = col_box.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.3)
        ctf.margin_top = Inches(0.3)
        cp = ctf.paragraphs[0]
        cp.text = col_title
        cp.font.size = Pt(16)
        cp.font.bold = True
        cp.font.color.rgb = col_color
        for t_name, t_desc in items:
            p_n = ctf.add_paragraph()
            p_n.text = f"• {t_name}"
            p_n.font.bold = True
            p_n.font.size = Pt(12)
            p_n.font.color.rgb = TEXT_WHITE
            p_n.space_before = Pt(8)
            p_d = ctf.add_paragraph()
            p_d.text = f"   {t_desc}"
            p_d.font.size = Pt(10.5)
            p_d.font.color.rgb = TEXT_MUTED
            p_d.space_before = Pt(1)

    # 6. Advantages & Disadvantages
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Comprehensive Evaluation: Advantages & Disadvantages", "Slide 06 • Critical Evaluation")

    adv_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.1))
    adv_box.fill.solid()
    adv_box.fill.fore_color.rgb = CARD_BG
    adv_box.line.color.rgb = BORDER_COL
    atf = adv_box.text_frame
    atf.word_wrap = True
    atf.margin_left = atf.margin_right = Inches(0.4)
    atf.margin_top = Inches(0.3)
    ap = atf.paragraphs[0]
    ap.text = "✅ Key Advantages"
    ap.font.size = Pt(18)
    ap.font.bold = True
    ap.font.color.rgb = SUCCESS_GREEN

    advs = [
        ("Unified Doorstep Journey", "Combines Rapido home pickup, intercity transport, and destination drop-off into a single booking flow."),
        ("Sub-Second AI Planning", "Gemini 2.5 Flash extracts all parameters instantly, eliminating tedious multi-form filtering."),
        ("Zero Lost-Screen Space UX", "Pinned header and footer with dedicated middle scrolling ensures controls never vanish."),
        ("Strict Itinerary Governance", "Immutable locked trips guarantee corporate policy compliance and prevent rogue expense inflation."),
        ("Zero Downtime Resilience", "Smart OSM fallback ensures seamless presentation even without live GPS connectivity.")
    ]
    for title, desc in advs:
        p_t = atf.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = SUCCESS_GREEN
        p_t.space_before = Pt(8)
        p_d = atf.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(1)

    dis_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.1))
    dis_box.fill.solid()
    dis_box.fill.fore_color.rgb = CARD_BG
    dis_box.line.color.rgb = BORDER_COL
    dtf = dis_box.text_frame
    dtf.word_wrap = True
    dtf.margin_left = dtf.margin_right = Inches(0.4)
    dtf.margin_top = Inches(0.3)
    dp = dtf.paragraphs[0]
    dp.text = "⚠️ Limitations & Future Work"
    dp.font.size = Pt(18)
    dp.font.bold = True
    dp.font.color.rgb = ROSE_RED

    diss = [
        ("External API Rate Limits", "Public Overpass & OSRM demo endpoints have throughput constraints, requiring local caching or paid geocoders for high scale."),
        ("Live Ride-Hailing Sandbox", "Current Rapido integration operates in simulated sandbox mode awaiting formal enterprise dispatch API access."),
        ("Locked State Inflexibility", "Post-booking lock prevents accidental edits, but handling emergency re-routes requires an admin override workflow."),
        ("LLM Prompt Hallucinations", "Ambiguous user inputs with unusual city spellings can occasionally require explicit fallback correction."),
        ("Multi-Currency Support", "Currently optimized for INR (₹) transactions; global corporate deployment will require multi-currency normalization.")
    ]
    for title, desc in diss:
        p_t = dtf.add_paragraph()
        p_t.text = f"✕ {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = ROSE_RED
        p_t.space_before = Pt(8)
        p_d = dtf.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(1)

    # 7. Conclusion Slide
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Conclusion & Future Roadmap", "Slide 07 • Project Impact")

    conc_banner = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.3))
    conc_banner.fill.solid()
    conc_banner.fill.fore_color.rgb = CARD_BG
    conc_banner.line.color.rgb = SUCCESS_GREEN
    c_tf = conc_banner.text_frame
    c_tf.word_wrap = True
    c_tf.margin_left = Inches(0.4)
    c_tf.margin_top = Inches(0.2)
    cp0 = c_tf.paragraphs[0]
    cp0.text = "🏁 EXECUTIVE SUMMARY"
    cp0.font.size = Pt(12)
    cp0.font.bold = True
    cp0.font.color.rgb = ACCENT_CYAN
    cp1 = c_tf.add_paragraph()
    cp1.text = "RoamAgent bridges the critical gap between fragmented travel portals and real-world door-to-door transit. By combining Google Gemini 2.5 Flash, real-time OSRM routing, OpenStreetMap discovery, and Rapido last-mile logistics, we deliver a frictionless, policy-compliant travel experience."
    cp1.font.size = Pt(13.5)
    cp1.font.color.rgb = TEXT_WHITE
    cp1.space_before = Pt(4)

    future_cards = [
        ("🚀 Production Dispatch API", "Transition from sandbox simulation to live enterprise dispatch APIs for Rapido, Uber, IRCTC, and airline GDS feeds for automated ticket issuance.", PRIMARY_BLUE, Inches(0.8)),
        ("⚡ Autonomous Re-routing", "Expand the Disruption Simulator to autonomously detect flight/train delays and automatically re-book connecting cabs and hotels in real time.", ACCENT_CYAN, Inches(4.8)),
        ("🌱 ESG & Carbon Optimization", "Implement smart route scoring prioritizing green corridor travel, electric bus fleets, and corporate carbon offsetting reports.", SUCCESS_GREEN, Inches(8.8))
    ]
    for title, desc, color, x_pos in future_cards:
        fc = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(3.1), Inches(3.7), Inches(3.7))
        fc.fill.solid()
        fc.fill.fore_color.rgb = CARD_BG
        fc.line.color.rgb = BORDER_COL
        ftf = fc.text_frame
        ftf.word_wrap = True
        ftf.margin_left = ftf.margin_right = Inches(0.3)
        ftf.margin_top = Inches(0.3)
        fp1 = ftf.paragraphs[0]
        fp1.text = title
        fp1.font.size = Pt(15)
        fp1.font.bold = True
        fp1.font.color.rgb = color
        fp2 = ftf.add_paragraph()
        fp2.text = desc
        fp2.font.size = Pt(11.5)
        fp2.font.color.rgb = TEXT_MUTED
        fp2.space_before = Pt(12)
        fp3 = ftf.add_paragraph()
        fp3.text = "Thank you! Questions & Live Demo."
        fp3.font.size = Pt(11)
        fp3.font.italic = True
        fp3.font.color.rgb = TEXT_WHITE
        fp3.space_before = Pt(20)

    output_path = r"c:\samartha hac\RoamAgent_Presentation_Deck.pptx"
    prs.save(output_path)
    print(f"Saved cleanly to: {output_path}")

if __name__ == "__main__":
    create_presentation()
